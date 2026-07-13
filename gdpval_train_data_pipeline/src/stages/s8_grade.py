"""
S8 · 打分 (跨家族 grader, 按 rubric + answer_key)
每份交付物由**与其 generator 跨家族**的 grader 打分; 记 path+分; 每 query 标 best。

产出: outputs/scores.jsonl
"""
from __future__ import annotations
from pathlib import Path

import util
import schema
from providers import base


def read_deliverable_text(path: str, max_chars=14000) -> str:
    if not path:
        return ""
    p = Path(path)
    if not p.exists():
        return ""
    suf = p.suffix.lower()
    try:
        if suf in (".md", ".html", ".htm", ".txt", ".csv"):
            return p.read_text(errors="ignore")[:max_chars]
        if suf == ".pdf":
            import fitz
            doc = fitz.open(str(p))
            t = "\n".join(doc[i].get_text() for i in range(doc.page_count))
            doc.close()
            return t[:max_chars]
        if suf == ".docx":
            import docx
            return "\n".join(par.text for par in docx.Document(str(p)).paragraphs)[:max_chars]
        if suf in (".xlsx", ".xlsm"):
            import openpyxl
            wb = openpyxl.load_workbook(str(p), data_only=True)
            out = []
            for ws in wb.worksheets:
                out.append(f"# sheet: {ws.title}")
                for row in ws.iter_rows(values_only=True):
                    cells = [str(c) for c in row if c is not None]
                    if cells:
                        out.append("\t".join(cells))
            return "\n".join(out)[:max_chars]
    except Exception as e:                      # noqa
        return f"[读取交付物失败 {suf}: {e}]"
    return f"[未支持格式 {suf}]"


def pick_graders(generator: str) -> list:
    """选与 generator **跨家族**的 grader (可多个)。pilot 只有 1 个跨族 (GLM↔DeepSeek);
    GPT/Gemini key 到位后自动变多 -> 触发多 grader 平均 + 分歧仲裁标记。"""
    graders = util.cfg("models")["roles"]["graders"]
    gfam = base.family_of(generator)
    cross = [g for g in graders if base.family_of(g) != gfam]
    if cross:
        return cross
    other = [g for g in graders if g != generator]
    return other or graders[:1]


def grade_one(grader, q, rubric, answer_key, deliverable_text) -> dict:
    sys_prompt = util.load_prompt("grader", q["domain"])
    rubric_lines = "\n".join(
        f'[{it["rubric_item_id"]}] (+{it["score"]}, {it["kind"]}'
        + (f'/{it["gdpval_dim"]}' if it.get("gdpval_dim") else "")
        + (", required" if it.get("required") else "")
        + f') {it["criterion"]}' for it in rubric["items"])
    ref = util.bundle_reference(q["reference_files"])
    user = (
        f"[任务]\n{q['prompt']}\n\n[format_spec]\n{q.get('format_spec','')}\n\n"
        f"[reference 数据]\n{ref}\n\n[answer_key 专家结论]\n{answer_key}\n\n"
        f"[rubric 条目]\n{rubric_lines}\n\n[待评交付物内容]\n{deliverable_text[:12000]}\n\n"
        f"逐条按 rubric_item_id 打分, 输出 per_item/dim_scores/total/justification。只 JSON。"
    )
    out = grader.chat_json(
        [{"role": "system", "content": sys_prompt},
         {"role": "user", "content": user}],
        schema=schema.GRADE_SCHEMA, temperature=0.0, max_tokens=4000)
    return out


def run(limit=None):
    rubrics = {r["query_id"]: r for r in util.read_jsonl(util.path("queries").parent / "rubrics.jsonl")}
    queries = {q["query_id"]: q for q in util.read_jsonl(util.path("queries").parent / "queries_passed.jsonl")}
    delivs = util.read_jsonl(util.path("queries").parent / "deliverables.jsonl")
    if limit:
        delivs = delivs[:limit]

    arb_thr = util.cfg("pipeline")["grading"].get("arbiter_on_disagreement", 0.25)
    workers = util.cfg("pipeline")["grading"].get("workers", 8)

    def _grade_deliv(d):
        qid = d["query_id"]
        q, rub = queries.get(qid), rubrics.get(qid)
        if not q or not rub:
            return None
        if not d["path"]:                             # agent 没产出文件 -> 0 分
            return schema.Score(deliverable_id=d["deliverable_id"], query_id=qid,
                                grader_model="-", total=0.0, max_score=rub["max_score"],
                                per_item={}, dim_scores={}).__dict__
        text = read_deliverable_text(d["path"])
        answer_key = util.bundle_answer_key(q["reference_files"])
        graded = []
        for gn in pick_graders(d["generator"])[:2]:   # 至多 2 个跨族 grader
            try:
                graded.append((gn, grade_one(base.get_provider(gn), q, rub, answer_key, text)))
            except Exception as e:                    # noqa
                print(f"  [S8] {d['deliverable_id']} grader={gn} 失败: {e}")
        if not graded:
            return None
        totals = [float(o.get("total", 0.0)) for _, o in graded]
        avg = sum(totals) / len(totals)
        arbitrated = (len(totals) >= 2 and
                      (max(totals) - min(totals)) / max(1, rub["max_score"]) > arb_thr)
        rep = graded[0][1]
        sc = schema.Score(
            deliverable_id=d["deliverable_id"], query_id=qid,
            grader_model="+".join(g for g, _ in graded), total=avg, max_score=rub["max_score"],
            per_item=rep.get("per_item", {}), dim_scores=rep.get("dim_scores", {}),
            arbitrated=arbitrated).__dict__
        print(f"  [S8] {d['deliverable_id']} ({d['generator']}->{sc['grader_model']}) "
              f"= {avg:.1f}/{rub['max_score']}" + (" ⚖分歧" if arbitrated else ""))
        return sc

    scores = [s for s in util.pmap(_grade_deliv, delivs, workers=workers) if s]
    by_query = {}
    for sc in scores:
        by_query.setdefault(sc["query_id"], []).append(sc)
    # 每 query 标 best
    for qid, lst in by_query.items():
        if lst:
            max(lst, key=lambda s: s["total"])["is_best"] = True

    util.write_jsonl(util.path("scores"), scores)
    print(f"[S8] 完成 {len(scores)} 份打分 -> {util.path('scores')}")
    return scores
